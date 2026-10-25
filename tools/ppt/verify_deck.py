"""Verificación estructural del deck: límites de slide y desbordes probables de texto."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[2]
PPTX = ROOT / "docs" / "presentacion" / "TP2-VoidHex-Tactics.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
TOL = 0.02

# ancho medio de carácter por punto, para Segoe UI / Consolas (aproximado)
CHAR_W = {"Segoe UI": 0.52, "Consolas": 0.62}


def inches(emu):
    return Emu(emu).inches


def main():
    deck = Presentation(PPTX)
    problems = 0
    warnings = 0

    print(f"slides: {len(deck.slides)}")
    print(f"size  : {inches(deck.slide_width):.2f} x {inches(deck.slide_height):.2f} in\n")

    for index, slide in enumerate(deck.slides, start=1):
        texts = []
        for shape in slide.shapes:
            x, y = inches(shape.left), inches(shape.top)
            w, h = inches(shape.width), inches(shape.height)

            if x < -TOL or y < -TOL or x + w > SLIDE_W + TOL or y + h > SLIDE_H + TOL:
                problems += 1
                print(f"  ! slide {index}: fuera de límites {shape.shape_type} "
                      f"({x:.2f},{y:.2f},{w:.2f},{h:.2f})")

            if not shape.has_text_frame:
                continue

            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    texts.append(run.text)

                    size = run.font.size.pt if run.font.size else 12
                    family = run.font.name or "Segoe UI"
                    per_char = CHAR_W.get(family, 0.52) * size / 72.0
                    needed = len(run.text) * per_char
                    if needed > w * 1.02 and "\n" not in run.text:
                        warnings += 1
                        print(f"  ~ slide {index}: posible desborde en {w:.2f}in "
                              f"({len(run.text)} chars @ {size:.0f}pt, necesita {needed:.2f}in)")
                        print(f"      {run.text[:70]!r}")

            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                texts.append(run.text)

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        first = texts[0][:58] if texts else "(sin texto)"
        print(f"  {index:02d}  {first}")
        if not notes:
            print(f"      ^ sin notas del orador")

    print(f"\nproblemas: {problems}   avisos: {warnings}")
    return problems


if __name__ == "__main__":
    raise SystemExit(main())
